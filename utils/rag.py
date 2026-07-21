"""RAG utility: loads and extracts text from the consulting firm's own
reference documents (service materials, past engagement summaries, etc.)."""

import os
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}
MAX_CONTEXT_CHARS = 60_000  # keep well within Claude's context window


def load_reference_documents(folder_path: str) -> str:
    """Return concatenated text extracted from all documents in *folder_path*."""
    if not folder_path or not os.path.exists(folder_path):
        return ""

    docs: list[str] = []
    for file_path in sorted(Path(folder_path).rglob("*")):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            text = _extract_text(file_path)
            if text and text.strip():
                docs.append(f"=== {file_path.name} ===\n{text.strip()}")

    if not docs:
        return ""

    combined = "\n\n".join(docs)
    if len(combined) > MAX_CONTEXT_CHARS:
        combined = combined[:MAX_CONTEXT_CHARS] + "\n\n[... 以下省略 ...]"

    return combined


def list_document_files(folder_path: str) -> list[Path]:
    """Return a list of supported document files found in *folder_path*."""
    if not folder_path or not os.path.exists(folder_path):
        return []
    return [
        p
        for p in sorted(Path(folder_path).rglob("*"))
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------

def _extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(file_path)
        elif suffix == ".docx":
            return _extract_docx(file_path)
        elif suffix == ".pptx":
            return _extract_pptx(file_path)
        elif suffix in {".txt", ".md"}:
            return file_path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001
        return f"[読み込みエラー: {exc}]"
    return ""


def _extract_pdf(file_path: Path) -> str:
    import pypdf  # type: ignore

    reader = pypdf.PdfReader(str(file_path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _extract_docx(file_path: Path) -> str:
    from docx import Document  # type: ignore

    doc = Document(str(file_path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(file_path: Path) -> str:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(file_path))
    parts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        if slide_texts:
            parts.append(f"[スライド {slide_num}]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)
